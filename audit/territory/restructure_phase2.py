#!/usr/bin/env python3
"""
Territory Restructure - Phase 2
Creates city territories, matches communities, updates all records, assigns FMs.
"""
import json, subprocess, csv, sys, os, re, time
import openpyxl

ORG = "dispatch"
OUTDIR = os.path.dirname(__file__)
OP_HOURS = "0OHVu000002ks2fOAA"

# ---- FM mapping ----
FM_MAP = {
    "Justin":  {"name": "Justin Smithwick",  "user_id": "005Vu00000J7hQHIAZ",  "resource_id": "0HnVu0000007mSvKAI"},
    "Trenton": {"name": "Trenton Kinard",    "user_id": "005Vu00000HRvzhIAD",  "resource_id": "0HnVu0000006ZhRKAU"},
    "Jamie":   {"name": "Jamie Hinson",      "user_id": "005Vu00000JHMZ3IAP",  "resource_id": "0HnVu00000081gUKAQ"},
    "Nick":    {"name": "Nick Melendez",     "user_id": "005Vu00000HRhdOIAT",  "resource_id": "0HnVu0000006dZlKAI"},
    "Scott":   {"name": "Scott Spaulding",   "user_id": "005Vu00000JHMXRIA5",  "resource_id": "0HnVu00000081gTKAQ"},
}

# ---- Zone IDs (created in phase 1) ----
ZONE_IDS = {
    1: "0HhVu0000001LSDKA2",  # North Zone + Asheville Extension
    2: "0HhVu0000001LTpKAM",  # Northeast Zone + Triad Extension  
    3: "0HhVu0000001LYfKAM",  # South/SE Zone + Columbia Extension
    4: "0HhVu0000001LVRKA2",  # West Zone + Greenville Extension
    5: "0HhVu0000001LX3KAM",  # Central Flex Zone
}

# ---- Extension city classification ----
EXTENSION_CITIES = {
    1: {"asheville", "hendersonville", "weaverville", "denver"},
    2: {"winston salem", "winston-salem", "kernersville", "kennersville", "king", "stokesdale", "archdale", "mocksville"},
    3: {"camden", "cayce", "gaston"},
    4: {"fountain inn", "travelers rest", "woodruff", "spartanburg", "boiling springs", "clinton"},
    5: set(),
}
EXTENSION_NAMES = {
    1: "Asheville Extension",
    2: "Triad Extension",
    3: "Columbia Extension",
    4: "Greenville Extension",
    5: None,
}

def run(cmd, silent=False):
    r = subprocess.run(cmd, capture_output=True, text=True)
    if not silent and r.returncode != 0:
        sys.stderr.write(f"WARN: {r.stderr[:200]}\n")
    return r.stdout, r.stderr, r.returncode

def sf_query(soql):
    cmd = ["sf", "data", "query", "--target-org", ORG, "--query", soql, "--json"]
    out, err, rc = run(cmd, silent=True)
    try:
        return json.loads(out).get("result", {}).get("records", [])
    except:
        return []

def sf_create(sobject, values_str):
    cmd = ["sf", "data", "create", "record", "--target-org", ORG,
           "--sobject", sobject, "--values", values_str, "--json"]
    out, err, rc = run(cmd, silent=True)
    try:
        d = json.loads(out)
        return d.get("result", {}).get("id")
    except:
        return None

def sf_delete(sobject, record_id):
    out, err, rc = run(["sf", "data", "delete", "record", "--target-org", ORG,
                        "--sobject", sobject, "--record-id", record_id, "--json"], silent=True)
    return rc == 0

def bulk_update_csv(sobject, rows):
    if not rows:
        return True
    csvpath = os.path.join(OUTDIR, f"_bulk_{sobject}.csv")
    with open(csvpath, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
    out, err, rc = run(["sf", "data", "update", "bulk", "--target-org", ORG,
                        "--sobject", sobject, "--file", csvpath, "--wait", "120", "--json"])
    try:
        d = json.loads(out)
        results = d.get("result", {})
        processed = results.get("numberRecordsProcessed", 0)
        failed = results.get("numberRecordsFailed", 0)
        print(f"    Bulk update {sobject}: {processed} processed, {failed} failed")
        return failed == 0
    except:
        print(f"    Bulk update parse error: {out[:300]}")
        return rc == 0

def normalize_name(n):
    """Normalize community name for matching."""
    n = str(n).lower().strip()
    # Normalize common variations
    n = re.sub(r'\s+', ' ', n)
    n = n.replace(' th', ' townhomes').replace(' ths', ' townhomes')
    n = re.sub(r'(\d+)s\b', r'\1', n)  # 60s -> 60
    n = n.replace("'", "").replace("-", " ")
    return n

def fuzzy_match(community_name, org_communities):
    """Find best match for a community name in org records."""
    norm_target = normalize_name(community_name)
    best_id = None
    best_score = 0
    
    for org_comm in org_communities:
        norm_org = normalize_name(org_comm["Name"])
        # Exact match
        if norm_target == norm_org:
            return org_comm["Id"], 100
        # Contains match
        if norm_target in norm_org or norm_org in norm_target:
            score = 80
        else:
            # Word overlap
            target_words = set(norm_target.split())
            org_words = set(norm_org.split())
            if len(target_words) > 0 and len(org_words) > 0:
                overlap = len(target_words & org_words) / max(len(target_words), len(org_words))
                score = int(overlap * 70)
            else:
                score = 0
        if score > best_score:
            best_score = score
            best_id = org_comm["Id"]
    
    return best_id if best_score >= 50 else None, best_score

# ============================================================
# Read spreadsheet
# ============================================================
print("=== Reading spreadsheet ===")
wb = openpyxl.load_workbook('/root/.claude/uploads/a160e003-5255-4957-a10a-33cca7cba422/4f53a66d-Zone_Rule_Breakdown.xlsx')
ws = wb['By Zone']

# zone_num -> set of cities
zone_cities = {1: set(), 2: set(), 3: set(), 4: set(), 5: set()}
spreadsheet_data = []

for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
    if not any(v is not None for v in row):
        continue
    zone_num, zone_name, community, city, zip_code, basis, fm, notes = (row + (None,)*8)[:8]
    if not zone_name or not community:
        continue
    zn = int(str(zone_num).replace("Zone ","").strip())
    city_str = str(city).strip() if city else ""
    city_norm = city_str.lower()
    
    # Determine territory name for this community
    if city_norm in EXTENSION_CITIES.get(zn, set()):
        terr_name = EXTENSION_NAMES.get(zn, city_str)
    else:
        terr_name = city_str if city_str else "Unknown"
    
    zone_cities[zn].add(city_str)
    spreadsheet_data.append({
        "zone_num": zn,
        "community_name": str(community).strip(),
        "city": city_str,
        "zip": str(zip_code).strip() if zip_code else "",
        "territory_name": terr_name,
        "fm": str(fm).strip() if fm else "",
    })

print(f"  Spreadsheet entries: {len(spreadsheet_data)}")

# ============================================================
# Create city-level child territories
# ============================================================
print("\n=== Creating city territories ===")

# Collect all unique territory names per zone
zone_territory_names = {}  # zone_num -> set of territory names
for entry in spreadsheet_data:
    zn = entry["zone_num"]
    tn = entry["territory_name"]
    if tn and tn != "Unknown":
        if zn not in zone_territory_names:
            zone_territory_names[zn] = set()
        zone_territory_names[zn].add(tn)

# Get existing child territories to avoid duplicates
existing_territories = sf_query("SELECT Id, Name, ParentTerritoryId FROM ServiceTerritory ORDER BY Name")
existing_by_name = {t["Name"]: t["Id"] for t in existing_territories}

city_territory_ids = {}  # (zone_num, territory_name) -> Id

for zone_num in sorted(zone_territory_names.keys()):
    parent_id = ZONE_IDS[zone_num]
    for terr_name in sorted(zone_territory_names[zone_num]):
        # Check if exists already with this parent
        key = (zone_num, terr_name)
        existing = next((t for t in existing_territories 
                         if t["Name"] == terr_name and t.get("ParentTerritoryId") == parent_id), None)
        if existing:
            city_territory_ids[key] = existing["Id"]
            print(f"  Exists: Zone{zone_num} '{terr_name}': {existing['Id']}")
            continue
        
        # Create it
        safe_name = terr_name.replace("'", "\\'")
        vals = f"Name='{safe_name}' IsActive=true OperatingHoursId={OP_HOURS} ParentTerritoryId={parent_id}"
        tid = sf_create("ServiceTerritory", vals)
        if tid:
            city_territory_ids[key] = tid
            print(f"  Created: Zone{zone_num} '{terr_name}': {tid}")
        else:
            # Maybe name already exists without right parent - find it
            if terr_name in existing_by_name:
                city_territory_ids[key] = existing_by_name[terr_name]
                print(f"  Found existing (diff parent): Zone{zone_num} '{terr_name}': {existing_by_name[terr_name]}")
            else:
                print(f"  FAILED: Zone{zone_num} '{terr_name}'")

print(f"  City territories created/found: {len(city_territory_ids)}")

# ============================================================
# Match spreadsheet communities to org communities
# ============================================================
print("\n=== Matching communities ===")
org_communities = sf_query(
    "SELECT Id, Name, Community_City__c, Community_Zip__c FROM Community__c ORDER BY Name"
)

matched = []
unmatched = []

for entry in spreadsheet_data:
    comm_name = entry["community_name"]
    zone_num = entry["zone_num"]
    
    # First try exact match
    org_match = next((c for c in org_communities if c["Name"].strip().lower() == comm_name.lower()), None)
    
    if org_match:
        matched.append({"entry": entry, "org_id": org_match["Id"], "org_name": org_match["Name"], "score": 100})
    else:
        # Fuzzy match
        best_id, score = fuzzy_match(comm_name, org_communities)
        if best_id:
            org_match = next(c for c in org_communities if c["Id"] == best_id)
            matched.append({"entry": entry, "org_id": best_id, "org_name": org_match["Name"], "score": score})
        else:
            unmatched.append(entry)

print(f"  Matched: {len(matched)}, Unmatched: {len(unmatched)}")
if unmatched:
    print("  Unmatched communities:")
    for u in unmatched:
        print(f"    Zone{u['zone_num']}: '{u['community_name']}' ({u['city']})")

# Save match report
with open(os.path.join(OUTDIR, "community_matches.json"), "w") as f:
    json.dump({"matched": matched, "unmatched": unmatched}, f, indent=2)

# ============================================================
# Update Community__c records
# ============================================================
print("\n=== Updating Community__c records ===")

update_rows = []
for m in matched:
    entry = m["entry"]
    zone_num = entry["zone_num"]
    zone_id = ZONE_IDS[zone_num]
    fm_key = entry["fm"]
    fm_data = FM_MAP.get(fm_key)
    
    row = {
        "Id": m["org_id"],
        "Zone__c": str(zone_num),
        "ServiceTerritory__c": zone_id,
        "Service_Territory_Name__c": entry["territory_name"],
    }
    if fm_data:
        row["FM_Assignment__c"] = fm_data["user_id"]
        row["Assigned_FM__c"] = fm_data["user_id"]
    
    update_rows.append(row)

print(f"  Updating {len(update_rows)} community records...")
if update_rows:
    # Write CSV for bulk update
    csvpath = os.path.join(OUTDIR, "_community_updates.csv")
    fieldnames = ["Id", "Zone__c", "ServiceTerritory__c", "Service_Territory_Name__c", "FM_Assignment__c", "Assigned_FM__c"]
    with open(csvpath, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for row in update_rows:
            writer.writerow({k: row.get(k, "") for k in fieldnames})
    
    out, err, rc = run(["sf", "data", "update", "bulk", "--target-org", ORG,
                        "--sobject", "Community__c", "--file", csvpath,
                        "--wait", "120", "--json"])
    try:
        d = json.loads(out)
        results = d.get("result", {})
        processed = results.get("numberRecordsProcessed", 0)
        failed = results.get("numberRecordsFailed", 0)
        print(f"  Community updates: {processed} processed, {failed} failed")
    except:
        print(f"  Update output: {out[:400]}")

# ============================================================
# Assign FMs to zone territories as ServiceTerritoryMember
# ============================================================
print("\n=== Assigning FMs to zone territories ===")

# Get existing members for new zones
existing_members = sf_query("SELECT Id, ServiceTerritoryId, ServiceResourceId FROM ServiceTerritoryMember")

for zone in [
    {"num": 1, "fm": "Justin"},
    {"num": 2, "fm": "Trenton"},
    {"num": 3, "fm": "Jamie"},
    {"num": 4, "fm": "Nick"},
    {"num": 5, "fm": "Scott"},
]:
    zone_id = ZONE_IDS[zone["num"]]
    fm_data = FM_MAP[zone["fm"]]
    resource_id = fm_data["resource_id"]
    
    # Check if already assigned
    existing = next((m for m in existing_members 
                     if m["ServiceTerritoryId"] == zone_id 
                     and m["ServiceResourceId"] == resource_id), None)
    if existing:
        print(f"  Zone {zone['num']}: {fm_data['name']} already assigned")
        continue
    
    vals = f"ServiceTerritoryId={zone_id} ServiceResourceId={resource_id} TerritoryType=P"
    mid = sf_create("ServiceTerritoryMember", vals)
    if mid:
        print(f"  Zone {zone['num']}: Assigned {fm_data['name']} as Primary member: {mid}")
    else:
        print(f"  Zone {zone['num']}: FAILED to assign {fm_data['name']}")

# ============================================================
# Delete old ServiceTerritoryMember records for old territories
# ============================================================
OLD_TERRITORY_IDS = [
    "0HhVu0000000ittKAA",  # Asheville
    "0HhVu00000016JcKAI",  # Cabarrus
    "0HhVu0000000isHKAQ",  # Charlotte - North
    "0HhVu00000016JaKAI",  # Charlotte - South
    "0HhVu00000016JZKAY",  # Charlotte Metro
    "0HhVu0000000ivVKAQ",  # Columbia
    "0HhVu00000016JfKAI",  # Foothills
    "0HhVu00000016JdKAI",  # Gaston
    "0HhVu0000000ix7KAA",  # Greenville
    "0HhVu00000016JbKAI",  # Lake Norman
    "0HhVu00000016JYKAY",  # Midlands SC
    "0HhVu0000001AoXKAU",  # Probe Terr
    "0HhVu00000016JiKAI",  # Spartanburg
    "0HhVu00000016JgKAI",  # Triad Metro
    "0HhVu00000016JhKAI",  # Triangle Metro
    "0HhVu00000016JXKAY",  # Upstate SC
    "0HhVu00000016JeKAI",  # York Lancaster
]

print("\n=== Removing members from old territories ===")
old_members = [m for m in existing_members if m["ServiceTerritoryId"] in OLD_TERRITORY_IDS]
print(f"  Old territory members to remove: {len(old_members)}")

for member in old_members:
    ok = sf_delete("ServiceTerritoryMember", member["Id"])
    status = "OK" if ok else "FAIL"
    print(f"  Delete member {member['Id']}: {status}")

# ============================================================
# Delete old ServiceTerritory records
# ============================================================
print("\n=== Deleting old service territories ===")

# First make sure communities are no longer pointing to old territories
# (we've already updated matched ones - unmatched ones may still point to old territories)
# For unmatched communities, clear their ServiceTerritory__c so they don't block deletion

# Check for any communities still pointing to old territories
still_pointing = sf_query(
    "SELECT Id, Name FROM Community__c WHERE ServiceTerritory__c IN " +
    "('" + "','".join(OLD_TERRITORY_IDS) + "')"
)
if still_pointing:
    print(f"  {len(still_pointing)} communities still pointing to old territories - clearing...")
    clear_rows = [{"Id": c["Id"], "ServiceTerritory__c": ""} for c in still_pointing]
    csvpath = os.path.join(OUTDIR, "_community_clear_st.csv")
    with open(csvpath, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["Id", "ServiceTerritory__c"])
        writer.writeheader()
        writer.writerows(clear_rows)
    out, err, rc = run(["sf", "data", "update", "bulk", "--target-org", ORG,
                        "--sobject", "Community__c", "--file", csvpath,
                        "--wait", "120", "--json"])
    print(f"  Clear result: {out[:200]}")

# Now delete old territories
for old_id in OLD_TERRITORY_IDS:
    # Check if it still has children
    children = sf_query(f"SELECT Id FROM ServiceTerritory WHERE ParentTerritoryId='{old_id}'")
    if children:
        print(f"  Skipping {old_id} (has {len(children)} children)")
        continue
    ok = sf_delete("ServiceTerritory", old_id)
    print(f"  Delete territory {old_id}: {'OK' if ok else 'FAIL'}")

# ============================================================
# Summary
# ============================================================
print("\n=== Summary ===")
final_territories = sf_query("SELECT Id, Name, ParentTerritoryId FROM ServiceTerritory ORDER BY Name")
print(f"  Final territory count: {len(final_territories)}")
for t in final_territories:
    parent_name = next((x["Name"] for x in final_territories if x["Id"] == t.get("ParentTerritoryId")), "(none)")
    print(f"    {t['Name']} (parent: {parent_name})")

final_members = sf_query(
    "SELECT ServiceTerritory.Name, ServiceResource.Name, TerritoryType "
    "FROM ServiceTerritoryMember ORDER BY ServiceTerritory.Name, ServiceResource.Name"
)
print(f"\n  Territory members: {len(final_members)}")
for m in final_members:
    print(f"    {m['ServiceTerritory']['Name']} | {m['ServiceResource']['Name']} | {m['TerritoryType']}")

print("\nDone.")
