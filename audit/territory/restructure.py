#!/usr/bin/env python3
"""
Territory Restructure Script
Creates new zone-based ServiceTerritory hierarchy, updates Communities, assigns FMs.
"""
import json, subprocess, csv, sys, os, re, time

ORG = "dispatch"
OUTDIR = os.path.dirname(__file__)

# ---- FM mapping from spreadsheet + org query ----
FM_MAP = {
    "Justin":  {"name": "Justin Smithwick",  "user_id": "005Vu00000J7hQHIAZ",  "resource_id": "0HnVu0000007mSvKAI"},
    "Trenton": {"name": "Trenton Kinard",    "user_id": "005Vu00000HRvzhIAD",  "resource_id": "0HnVu0000006ZhRKAU"},
    "Jamie":   {"name": "Jamie Hinson",      "user_id": "005Vu00000JHMZ3IAP",  "resource_id": "0HnVu00000081gUKAQ"},
    "Nick":    {"name": "Nick Melendez",     "user_id": "005Vu00000HRhdOIAT",  "resource_id": "0HnVu0000006dZlKAI"},
    "Scott":   {"name": "Scott Spaulding",   "user_id": "005Vu00000JHMXRIA5",  "resource_id": "0HnVu00000081gTKAQ"},
}
JERSAIN = {"name": "Jersain Laris", "user_id": "005Vu00000JHMfVIAX", "resource_id": "0HnVu00000081gVKAQ"}

# ---- Zone definitions ----
ZONES = [
    {"num": 1, "name": "North Zone + Asheville Extension",         "fm": "Justin"},
    {"num": 2, "name": "Northeast Zone + Triad Extension",          "fm": "Trenton"},
    {"num": 3, "name": "South / Southeast Zone + Columbia Extension","fm": "Jamie"},
    {"num": 4, "name": "West Zone + Greenville Extension",          "fm": "Nick"},
    {"num": 5, "name": "Central Flex Zone",                         "fm": "Scott"},
]

# Extension cities per zone - these communities get the extension territory name
EXTENSION_CITIES = {
    1: {"Asheville", "Hendersonville", "Weaverville", "Denver"},
    2: {"Winston Salem", "Winston-Salem", "Kernersville", "Kennersville", "King", "Stokesdale", "Archdale", "Mocksville"},
    3: {"Camden", "Cayce", "Gaston"},
    4: {"Fountain Inn", "Travelers Rest", "Woodruff", "Spartanburg", "Boiling Springs", "Clinton"},
    5: set(),
}
EXTENSION_NAMES = {
    1: "Asheville Extension",
    2: "Triad Extension",
    3: "Columbia Extension",
    4: "Greenville Extension",
    5: None,
}

def run(cmd, silent=False):
    r = subprocess.run(cmd, capture_output=True, text=True)
    if not silent and r.returncode != 0:
        print(f"  WARN: {' '.join(cmd[:4])} stderr: {r.stderr[:200]}")
    return r.stdout, r.stderr, r.returncode

def sf_query(soql, tooling=False):
    cmd = ["sf", "data", "query", "--target-org", ORG, "--query", soql, "--json"]
    if tooling:
        cmd.insert(3, "--use-tooling-api")
    out, err, rc = run(cmd, silent=True)
    try:
        d = json.loads(out)
        return d.get("result", {}).get("records", [])
    except:
        return []

def sf_create(sobject, fields):
    """Create a single record, return Id."""
    args = []
    for k, v in fields.items():
        args += ["--values", f"{k}={v}"]
    out, err, rc = run(["sf", "data", "create", "record",
                        "--target-org", ORG, "--sobject", sobject] + args)
    if rc == 0:
        try:
            d = json.loads(out)
            rid = d.get("result", {}).get("id", "")
            if rid:
                return rid
        except:
            pass
    # try json flag
    out, err, rc = run(["sf", "data", "create", "record",
                        "--target-org", ORG, "--sobject", sobject, "--json"] + args)
    try:
        d = json.loads(out)
        rid = d.get("result", {}).get("id", "")
        return rid
    except:
        print(f"  CREATE FAIL {sobject}: {err[:200]}")
        return None

def sf_update(sobject, record_id, fields):
    args = ["--record-id", record_id]
    for k, v in fields.items():
        if v is None:
            continue
        args += ["--values", f"{k}={v}"]
    out, err, rc = run(["sf", "data", "update", "record",
                        "--target-org", ORG, "--sobject", sobject, "--json"] + args)
    return rc == 0

def sf_delete(sobject, record_id):
    out, err, rc = run(["sf", "data", "delete", "record",
                        "--target-org", ORG, "--sobject", sobject,
                        "--record-id", record_id, "--json"])
    return rc == 0

def bulk_upsert_csv(sobject, rows, ext_id_field=None):
    """Write CSV and bulk upsert."""
    if not rows:
        return True
    csvpath = os.path.join(OUTDIR, f"_bulk_{sobject}.csv")
    with open(csvpath, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
    cmd = ["sf", "data", "upsert", "bulk", "--target-org", ORG,
           "--sobject", sobject, "--file", csvpath,
           "--external-id", ext_id_field or "Id", "--wait", "60", "--json"]
    out, err, rc = run(cmd)
    try:
        d = json.loads(out)
        results = d.get("result", {})
        print(f"    bulk upsert {sobject}: {results.get('numberRecordsProcessed',0)} processed, {results.get('numberRecordsFailed',0)} failed")
        return results.get("numberRecordsFailed", 1) == 0
    except:
        print(f"    bulk upsert output: {out[:300]} {err[:200]}")
        return rc == 0


# ============================================================
# STEP 1: Query current state
# ============================================================
print("\n=== STEP 1: Query current state ===")
existing_territories = sf_query("SELECT Id, Name, ParentTerritoryId FROM ServiceTerritory ORDER BY Name")
print(f"  Existing territories: {len(existing_territories)}")
for t in existing_territories:
    print(f"    {t['Id']} | {t['Name']}")

existing_members = sf_query("SELECT Id, ServiceTerritoryId, ServiceResourceId FROM ServiceTerritoryMember")
print(f"  Existing territory members: {len(existing_members)}")

communities = sf_query(
    "SELECT Id, Name, Community_City__c, Community_Zip__c, Zone__c, "
    "ServiceTerritory__c, FM_Assignment__c, Assigned_FM__c FROM Community__c ORDER BY Name"
)
print(f"  Communities: {len(communities)}")

# ============================================================
# STEP 2: Read spreadsheet data
# ============================================================
print("\n=== STEP 2: Read spreadsheet ===")
import openpyxl
wb = openpyxl.load_workbook('/root/.claude/uploads/a160e003-5255-4957-a10a-33cca7cba422/4f53a66d-Zone_Rule_Breakdown.xlsx')
ws = wb['By Zone']

spreadsheet_communities = []
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
    if not any(v is not None for v in row):
        continue
    zone_num, zone_name, community, city, zip_code, basis, fm, notes = (row + (None,)*8)[:8]
    if zone_name and community:
        zone_num_int = int(zone_num.replace("Zone ","").strip()) if zone_num else None
        spreadsheet_communities.append({
            "zone_num": zone_num_int,
            "zone_name": zone_name,
            "community": str(community).strip(),
            "city": str(city).strip() if city else "",
            "zip": str(zip_code).strip() if zip_code else "",
            "basis": basis,
            "fm": fm,
        })
print(f"  Spreadsheet communities: {len(spreadsheet_communities)}")

# ============================================================
# STEP 3: Create zone parent territories
# ============================================================
print("\n=== STEP 3: Create zone parent territories ===")

# Get OperatingHours ID (required for ServiceTerritory)
op_hours = sf_query("SELECT Id, Name FROM OperatingHours ORDER BY Name LIMIT 5")
op_hours_id = op_hours[0]["Id"] if op_hours else None
print(f"  Operating hours: {[(o['Id'], o['Name']) for o in op_hours]}")

zone_territory_ids = {}  # zone_num -> territory Id

for zone in ZONES:
    # Check if already exists
    existing = next((t for t in existing_territories if t["Name"] == zone["name"]), None)
    if existing:
        print(f"  Zone {zone['num']} territory already exists: {existing['Id']}")
        zone_territory_ids[zone["num"]] = existing["Id"]
        continue
    
    fields = {"Name": zone["name"], "IsActive": "true"}
    if op_hours_id:
        fields["OperatingHoursId"] = op_hours_id
    
    tid = sf_create("ServiceTerritory", fields)
    if tid:
        zone_territory_ids[zone["num"]] = tid
        print(f"  Created Zone {zone['num']} '{zone['name']}': {tid}")
    else:
        print(f"  FAILED to create Zone {zone['num']}")

print(f"  Zone territory IDs: {zone_territory_ids}")

# Save to file for reference
with open(os.path.join(OUTDIR, "zone_territory_ids.json"), "w") as f:
    json.dump(zone_territory_ids, f, indent=2)

